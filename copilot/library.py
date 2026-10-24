"""Knowledge library: indexes a folder of documents and searches it with BM25."""
import math
import re
import shutil
import time
from collections import Counter
from pathlib import Path

TEXT_EXT = {".md", ".txt", ".csv", ".json", ".html", ".htm", ".rtf", ".tex", ".yaml", ".yml"}
DOC_EXT = TEXT_EXT | {".pdf", ".docx", ".pptx"}
CHUNK = 900
OVERLAP = 150
MIRROR = Path.home() / ".cache" / "meeting-copilot" / "library-text"
STOP = set("a an the and or of to in on for with is are was were be been it this that as at by from we you i they our your not but if so do does did have has had can will would should could about into than then there their what which who how when where why".split())


def tokens(text: str) -> list[str]:
    return [w for w in re.findall(r"[a-z0-9][a-z0-9'\-\.]*[a-z0-9]|[a-z0-9]", text.lower()) if w not in STOP]


def extract(path: Path) -> str:
    ext = path.suffix.lower()
    try:
        if ext in TEXT_EXT:
            text = path.read_text(errors="ignore")
            if ext in {".html", ".htm"}:
                text = re.sub(r"<script.*?</script>|<style.*?</style>", " ", text, flags=re.S | re.I)
                text = re.sub(r"<[^>]+>", " ", text)
            return text
        if ext == ".pdf":
            from pypdf import PdfReader
            return "\n".join(f"[page {i + 1}] " + (p.extract_text() or "") for i, p in enumerate(PdfReader(str(path)).pages))
        if ext == ".docx":
            import docx
            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                parts += [" | ".join(c.text for c in row.cells) for row in t.rows]
            return "\n".join(parts)
        if ext == ".pptx":
            from pptx import Presentation
            out = []
            for i, slide in enumerate(Presentation(str(path)).slides):
                out.append(f"[slide {i + 1}]")
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        out.append(shape.text_frame.text)
            return "\n".join(out)
    except Exception as e:  # unreadable files are skipped, not fatal
        print(f"library: could not read {path.name}: {e}")
    return ""


class Library:
    def __init__(self, folder: str | None):
        self.folder = Path(folder).expanduser() if folder else None
        self.chunks: list[dict] = []
        self.texts: dict[str, str] = {}
        self.mtimes: dict[str, float] = {}
        self.df: Counter = Counter()
        self.avgdl = 1.0
        self.indexed_at = 0.0
        self.mirror = MIRROR

    def files(self) -> list[Path]:
        if not self.folder or not self.folder.is_dir():
            return []
        return sorted(p for p in self.folder.rglob("*") if p.is_file() and p.suffix.lower() in DOC_EXT and not p.name.startswith((".", "~$")))

    def refresh(self) -> bool:
        """Re-index when files were added, removed or changed. Returns True if it re-indexed."""
        files = self.files()
        mt = {str(p): p.stat().st_mtime for p in files}
        if mt == self.mtimes:
            return False
        texts = {}
        for p in files:
            key = str(p)
            texts[key] = self.texts[key] if self.mtimes.get(key) == mt[key] and key in self.texts else extract(p)
        self.texts, self.mtimes = texts, mt
        self.write_mirror()
        self.chunks = []
        for key, text in texts.items():
            text = re.sub(r"[ \t]+", " ", text)
            for start in range(0, max(len(text), 1), CHUNK - OVERLAP):
                piece = text[start:start + CHUNK].strip()
                if len(piece) < 40:
                    continue
                toks = tokens(piece + " " + Path(key).stem.replace("_", " ").replace("-", " "))
                self.chunks.append({"file": key, "start": start, "text": piece, "tf": Counter(toks), "len": len(toks)})
        self.df = Counter()
        for c in self.chunks:
            self.df.update(c["tf"].keys())
        self.avgdl = sum(c["len"] for c in self.chunks) / max(len(self.chunks), 1)
        self.indexed_at = time.time()
        return True

    def write_mirror(self):
        """Plain-text copies of every document, so Claude's Grep and Read work on PDFs and Office files too."""
        shutil.rmtree(self.mirror, ignore_errors=True)
        self.mirror.mkdir(parents=True, exist_ok=True)
        for key, text in self.texts.items():
            out = self.mirror / (self.rel(key) + ".txt")
            out.parent.mkdir(parents=True, exist_ok=True)
            out.write_text(text)

    def rel(self, key: str) -> str:
        try:
            return str(Path(key).relative_to(self.folder))
        except Exception:
            return Path(key).name

    def search(self, query: str, k: int = 5) -> list[dict]:
        q = tokens(query)
        if not q or not self.chunks:
            return []
        n = len(self.chunks)
        scored = []
        for c in self.chunks:
            s = 0.0
            for t in q:
                f = c["tf"].get(t)
                if not f:
                    continue
                idf = math.log(1 + (n - self.df[t] + 0.5) / (self.df[t] + 0.5))
                s += idf * f * 2.2 / (f + 1.2 * (0.25 + 0.75 * c["len"] / self.avgdl))
            if s > 0:
                scored.append((s, c))
        scored.sort(key=lambda x: -x[0])
        out, seen = [], set()
        for s, c in scored:
            key = (c["file"], c["start"] // CHUNK)
            if key in seen:
                continue
            seen.add(key)
            out.append({"file": self.rel(c["file"]), "score": round(s, 2), "text": c["text"]})
            if len(out) >= k:
                break
        return out

    def read(self, name: str, offset: int = 0, length: int = 8000) -> str | None:
        for key, text in self.texts.items():
            if self.rel(key) == name or Path(key).name == name:
                return text[offset:offset + length]
        return None

    def summary(self) -> dict:
        return {"folder": str(self.folder) if self.folder else "", "files": len(self.texts), "chunks": len(self.chunks)}
